"""Load leads from a Notion database export (CSV, or the same opened and saved in Excel as .xlsx).

Columns are matched by name, ignoring case, so the export works as Notion writes it: Email, Contact name,
Business, Entity type, Status, No website and Lead ID. Only Email is required.
"""

from __future__ import annotations

import csv
import io
from collections import Counter
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import validate_email

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
#: The ICO treats these as corporate subscribers; sole traders and partnerships are individuals.
CORPORATE = ("limited company", "llp", "plc")

COLUMNS = {
    "email": ("email", "e-mail", "email address"),
    "name": ("contact name", "first name", "contact", "name"),
    "business": ("business", "business name", "company", "company name"),
    "entity_type": ("entity type",),
    "status": ("status",),
    "no_website": ("no website",),
    "lead_ref": ("lead id", "id"),
}
TRUE = {"yes", "true", "1", "checked", "x", "✓", "☑"}
TITLES = {"mr", "mrs", "ms", "miss", "mx", "dr", "prof", "sir", "dame"}


def first_name_from(contact: str) -> str:
    """The first name to greet, from Notion's "Contact name" ("Helen Claire Allen (director)").

    Several people ("Jamie Ireland, Elizabeth Ireland (directors)") give no name: the address is likely
    shared, so the email opens "Hi there,". Titles are dropped, and an all-capitals name is title-cased.
    """
    people = contact.split("(", 1)[0]
    if "," in people or " and " in people.lower() or "&" in people:
        return ""
    words = [word for word in people.split() if word.lower().strip(".") not in TITLES]
    if not words or not any(char.isalpha() for char in words[0]):
        return ""
    first = words[0].strip(".,")
    return (first.title() if first.isupper() else first)[:60]


@dataclass
class Row:
    email: str
    first_name: str
    business: str
    entity_type: str
    lead_ref: str


@dataclass
class Parsed:
    rows: list[Row] = field(default_factory=list)
    total: int = 0
    #: Rows left out, by reason (shown after the import).
    skipped: Counter = field(default_factory=Counter)


def _text(value) -> str:
    return " ".join(str(value if value is not None else "").split())


def _table(upload) -> list[dict[str, str]]:
    name = (getattr(upload, "name", "") or "").lower()
    if getattr(upload, "size", 0) > MAX_BYTES:
        raise ValidationError("The file is larger than 5 MB. Export fewer rows, or split it.")
    data = upload.read(MAX_BYTES + 1)
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
        reader = csv.reader(io.StringIO(text))
        lines = list(reader)
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook  # parses with defusedxml when it is installed (it is)

        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - any unreadable workbook gets the same answer
            raise ValidationError("This .xlsx file could not be read. Save it again from Excel.") from None
        sheet = book.worksheets[0]
        lines = [[_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        book.close()
    else:
        raise ValidationError("Upload the Notion export as .csv, or as .xlsx saved from Excel.")
    if not lines:
        raise ValidationError("The file is empty.")
    headers = [_text(cell).lower() for cell in lines[0]]
    if len(lines) - 1 > MAX_ROWS:
        raise ValidationError(f"The file has more than {MAX_ROWS} rows. Split it and load each part.")
    return [dict(zip(headers, (_text(cell) for cell in line), strict=False)) for line in lines[1:]]


def _column(headers: list[str], key: str) -> str | None:
    for name in COLUMNS[key]:
        if name in headers:
            return name
    return None


def parse(upload, *, only_new: bool, corporate_only: bool, no_website_only: bool) -> Parsed:
    """Rows to queue and rows left out. Filters apply only where the export has the column, except
    corporate_only, which needs an Entity type column to tell who may be emailed."""
    table = _table(upload)
    headers = list(table[0]) if table else []
    cols = {key: _column(headers, key) for key in COLUMNS}
    if cols["email"] is None:
        raise ValidationError("The file has no Email column. Include the Email property in the export.")
    if corporate_only and cols["entity_type"] is None:
        raise ValidationError(
            "The file has no Entity type column, so companies cannot be told from sole traders. Include "
            "Entity type in the export, or untick “Limited companies, LLPs and PLCs only”."
        )
    parsed = Parsed(total=len(table))
    seen: set[str] = set()
    for record in table:
        if not any(record.values()):
            parsed.total -= 1
            continue

        def get(key, record=record):
            return record.get(cols[key], "") if cols[key] else ""

        email = get("email").strip().lower().removeprefix("mailto:")
        if not email:
            parsed.skipped["No email address"] += 1
            continue
        try:
            validate_email(email)
        except ValidationError:
            parsed.skipped["Not a valid email address"] += 1
            continue
        if only_new and cols["status"] and get("status").lower() != "new":
            parsed.skipped["Status is not New"] += 1
            continue
        if corporate_only and get("entity_type").lower() not in CORPORATE:
            parsed.skipped["Not a limited company, LLP or PLC"] += 1
            continue
        if no_website_only and cols["no_website"] and get("no_website").lower() not in TRUE:
            parsed.skipped["Has a website"] += 1
            continue
        if email in seen:
            parsed.skipped["Listed twice in the file"] += 1
            continue
        seen.add(email)
        parsed.rows.append(
            Row(
                email=email[:254],
                first_name=first_name_from(get("name")),
                business=get("business")[:200],
                entity_type=get("entity_type")[:40],
                lead_ref=get("lead_ref")[:40],
            )
        )
    return parsed
